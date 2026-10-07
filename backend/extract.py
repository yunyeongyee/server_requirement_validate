# PDF / DOCX / XLSX 요구사항 추출

"""문서 본문 추출 + 서버 요구사항 규칙 기반 추출."""
from __future__ import annotations
import csv, io, json, re, uuid
from pathlib import Path


# ------------------------------------------------------------------ 본문 추출
def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        cands = []
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(data)
            cands.append("\n".join(pdf[i].get_textpage().get_text_range().replace("\r\n", "\n").replace("\r", "\n") for i in range(len(pdf))))
            pdf.close()
        except Exception:
            pass
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            cands.append("\n".join((p.extract_text() or "") for p in reader.pages))
        except Exception:
            pass
        return max(cands, key=_text_quality) if cands else ""
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    text = _decode(data)
    if ext == ".csv":
        return "\n".join(" | ".join(r) for r in csv.reader(io.StringIO(text)) if any(r))
    if ext == ".json":
        try:
            return "\n".join(_flatten_json(json.loads(text)))
        except json.JSONDecodeError:
            return text
    if ext == ".txt":
        return text
    raise ValueError(f"지원하지 않는 형식: {ext}")


def extract_document_context(filename: str, data: bytes) -> dict:
    """표와 시트 경계를 유지한 AI 분석용 문서 맥락을 만든다."""
    ext = Path(filename).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheets = []
        for worksheet in workbook.worksheets:
            rows = []
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                cells = [str(value).strip() if value is not None else "" for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append({"number": row_number, "cells": cells})
            sheets.append({"name": worksheet.title, "rows": rows})
        workbook.close()
        return {"format": "xlsx", "sheets": sheets}
    if ext == ".docx":
        import docx
        document = docx.Document(io.BytesIO(data))
        tables = []
        for index, table in enumerate(document.tables, 1):
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
            if rows:
                tables.append({"number": index, "rows": rows})
        return {
            "format": "docx",
            "paragraphs": [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()],
            "tables": tables,
        }
    if ext == ".pdf":
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(data)
            pages = [
                {"number": index + 1, "text": pdf[index].get_textpage().get_text_range().strip()}
                for index in range(len(pdf))
            ]
            pdf.close()
        except Exception:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            pages = [{"number": index + 1, "text": (page.extract_text() or "").strip()}
                     for index, page in enumerate(reader.pages)]
        return {"format": "pdf", "pages": [page for page in pages if page["text"]]}
    if ext == ".csv":
        text = _decode(data)
        rows = [{"number": index, "cells": [cell.strip() for cell in row]}
                for index, row in enumerate(csv.reader(io.StringIO(text)), 1) if any(cell.strip() for cell in row)]
        return {"format": "csv", "rows": rows}
    return {"format": ext.lstrip(".") or "text", "text": extract_text(filename, data)}


def _text_quality(t: str) -> float:
    """읽을 수 있는 한글/영숫자 비율 + 서버 관련 키워드 수."""
    if not t.strip():
        return 0
    good = sum(1 for ch in t if ("가" <= ch <= "힣") or ch.isascii() and (ch.isalnum() or ch in " .,:()-/%"))
    kw = len(re.findall(r"cpu|메모리|memory|프로세서|전원|디스크|nic|raid|psu|ssd|hdd", t, re.I))
    return good / len(t) + kw * 0.02 - t.count("(cid:") * 0.01


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp949", "euc-kr", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "ignore")


def _flatten_json(obj, prefix="") -> list[str]:
    out = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            out += _flatten_json(v, f"{prefix}{k}: " if not isinstance(v, (dict, list)) else f"{prefix}{k} > ")
    elif isinstance(obj, list):
        for v in obj:
            out += _flatten_json(v, prefix)
    else:
        out.append(f"{prefix}{obj}")
    return out


# ------------------------------------------------------------------ 요구사항 추출
# key 정의: 검증기가 이해하는 표준 키
KEYS = {
    "memory_gb":     ("Memory", "GB"),
    "cpu_sockets":   ("CPU", "Socket"),
    "nic_speed_gb":  ("NIC Speed", "GbE"),
    "nic_ports":     ("NIC Port", "Port"),
    "fc_speed_gb":   ("FC Speed", "Gb"),
    "fc_ports":      ("FC Port", "Port"),
    "ocp_required":  ("OCP 3.0", ""),
    "raid_level":    ("RAID", ""),
    "dual_psu":      ("Dual PSU", ""),
    "psu_watt":      ("PSU Capacity", "W"),
    "free_pcie":     ("Free PCIe Slot", "EA"),
    "gpu_count":     ("GPU", "EA"),
    "cpu_cores":     ("CPU Core", "Core"),
    "disk_count":    ("Disk", "EA"),
    "disk_size_gb":  ("Disk Size", "GB"),
    "disk_total_gb": ("Disk Total", "GB"),
    "rack_mount":    ("Rack Type", ""),
    "raid_controller": ("RAID Controller", ""),
}

VAGUE = re.compile(r"(충분한|충분히|적절한|적정|안정적|고성능|최적|원활|유연한|확장성|우수한|향후\s*고려|등\s*고려|협의)")
NUM = r"(\d+(?:\.\d+)?)"
GE = r"(이상|以上|or more|minimum|min\.?|at least|\+)"


def _req(key, op, value, line, note="", status="auto", **extra):
    label, unit = KEYS[key]
    return {"id": uuid.uuid4().hex[:8], "key": key, "label": label, "op": op,
            "value": value, "unit": unit, "source": line.strip()[:200],
            "status": status, "note": note, **extra}


def _op(line):
    """용량·수량 기준은 '정확히/만/only'가 있을 때만 같음(=), 그 외는 이상(>=).
    요구사항의 '메모리 128GB'는 보통 최소 기준이라 더 큰 구성을 미충족으로 보면 안 된다."""
    return "=" if re.search(r"정확히|exactly|\bonly\b|\d\s*(?:gb|tb|w|개|ea|소켓|socket|core|코어)?\s*만\b", line, re.I) else ">="


def _to_gb(n: float, unit: str) -> float:
    return n * 1024 if unit.upper().startswith("T") else n


def normalize(text: str) -> str:
    """PDF 추출 텍스트 정리: 전각→반각, '메 모 리' 처럼 글자마다 띄어진 한글 붙이기, 공백 정리."""
    import unicodedata
    t = unicodedata.normalize("NFKC", text).replace("\u00a0", " ").replace("\t", " ")
    out = []
    for line in t.splitlines():
        toks = line.split()
        if len(toks) >= 4 and sum(len(x) == 1 for x in toks) / len(toks) > 0.6:
            line = re.sub(r"(?<=\S) (?=\S)", "", line)       # 글자 단위로 띄어진 줄
        line = re.sub(r"(\d)\s+(GB|TB|GbE|Gb|G|W|Port|포트|개|EA)\b", r"\1\2", line, flags=re.I)
        line = re.sub(r"이\s+상", "이상", line)
        out.append(re.sub(r"\s{2,}", " ", line).strip())
    return "\n".join(out)


def _segments(text: str) -> list[str]:
    """문장 단위 + PDF 줄바꿈으로 끊긴 항목을 위해 '항목명 줄 + 값 줄' 결합본도 함께 검사."""
    lines = [l for l in re.split(r"\n|。|(?<=[.;])\s+", text) if l.strip()]
    segs = list(lines)
    for a, b in zip(lines, lines[1:]):
        if len(a) < 40 and not re.search(r"\d", a) and re.search(r"\d", b):
            segs.append(a + " " + b)
    return segs


def _ports(L: str) -> int | None:
    """'총 4포트' 우선, 아니면 '2Port x 2EA' → 4."""
    t = re.search(r"(?:총|total|합계)\s*(\d+)\s*(?:port|포트|p\b)", L)
    if t:
        return int(t.group(1))
    p = re.search(NUM + r"\s*(?:port|포트|p\b)", L)
    word = re.search(r"(single|dual|quad|octa|싱글|듀얼|쿼드)\s*-?\s*(?:port|포트)", L)
    if not p and not word:
        return None
    if not p:
        # 'Dual Port × 2' → 2 × 2
        n = {"single": 1, "싱글": 1, "dual": 2, "듀얼": 2, "quad": 4, "쿼드": 4, "octa": 8}[word.group(1)]
        mult = re.search(r"(?:port|포트)[^x×*\d]{0,12}[x×*]\s*(\d+)\s*(?:ea|개|장|식)?", L[word.start():])
        return n * int(mult.group(1)) if mult else n
    n = int(float(p.group(1)))
    mult = re.search(r"(?:port|포트|p\b)[^x×*\d]{0,12}[x×*]\s*(\d+)\s*(?:ea|개|장|식)?", L[p.start():])
    return n * int(mult.group(1)) if mult else n


def extract_requirements(text: str) -> list[dict]:
    reqs: list[dict] = []
    for raw in _segments(normalize(text)):
        line = raw.strip(" -•*·\t")
        if len(line) < 3:
            continue
        L = line.lower()
        found = []

        # Memory
        if re.search(r"memory|메모리|\bram\b|램|dimm|주\s*기억\s*장치|기억\s*용량", L):
            tot = re.search(r"(?:총|total|합계|전체)\s*[:(]?\s*" + NUM + r"\s*(gb|tb)", L)
            m = None if tot else re.search(NUM + r"\s*(gb|tb)[^\d,;]{0,25}?[x×*]\s*" + NUM, L)
            if tot:
                found.append(_req("memory_gb", ">=" if re.search(GE, L) else "=", _to_gb(float(tot.group(1)), tot.group(2)), line))
            elif m:
                found.append(_req("memory_gb", ">=" if re.search(GE, L) else "=",
                                  _to_gb(float(m.group(1)), m.group(2)) * float(m.group(3)), line))
            if not tot and not m:
                m = re.search(NUM + r"\s*(gb|tb)", L)
                if m:
                    found.append(_req("memory_gb", _op(L), _to_gb(float(m.group(1)), m.group(2)), line))

        # CPU sockets
        if re.search(r"cpu|프로세서|processor|소켓|socket|중앙\s*처리|xeon|epyc", L):
            m = re.search(r"[x×*]\s*(\d)\s*(?:ea|개|소켓|socket)?\b(?!\s*(?:ghz|core|코어|gb|tb|mb|w\b))|(?<![\w.])(\d)\s*(socket|소켓|ea|개|way|cpu|p\b|식)|(?:cpu|프로세서)\s*[x×*:]\s*(\d)\b(?!\s*(?:ghz|core|코어|gb|mb))|dual\s*(socket|cpu)|2\s*-?\s*way", L)
            if m:
                n = 2 if (m.group(5) or "dual" in L or "2-way" in L) else int(m.group(1) or m.group(2) or m.group(4))
                if 1 <= n <= 8:
                    found.append(_req("cpu_sockets", _op(L), n, line))

        # CPU 코어 수 ("코어 32개 이상", "32C", "32 core") — 'CPU당/소켓당'이 있으면 CPU 1개 기준
        m = re.search(r"(?:코어|core)\s*(?:수\s*)?[:：]?\s*(\d{1,3})\s*(?:개|ea)?|(\d{1,3})\s*(?:코어|cores?|c\b)(?!\s*(?:ghz|gb))", L)
        if m and re.search(r"cpu|프로세서|processor|코어|core|xeon|epyc|\d\s*c\b", L):
            n = int(m.group(1) or m.group(2))
            if 2 <= n <= 512:
                # 'Xeon Gold 6430 32C' 처럼 모델 표기 옆의 코어 수는 CPU 1개 기준
                per = bool(re.search(r"cpu\s*당|소켓\s*당|per\s*(?:cpu|socket|processor)|프로세서\s*당", L)
                           or (m.group(2) and re.search(r"xeon|epyc|ampere", L)))
                found.append(_req("cpu_cores", _op(L), n, line, note="CPU당" if per else "총 코어"))

        # 한 줄에 FC와 Ethernet이 같이 있으면 쉼표·슬래시로 나눠 각자 읽는다 ("10GbE 2포트, FC 32Gb 2포트")
        fc_re = r"\bfc\b|fibre|fiber channel|hba"
        clauses = [c for c in re.split(r"[,;/]|\s및\s|\s그리고\s", L) if c.strip()]
        fc_part = " ".join(c for c in clauses if re.search(fc_re, c) and not re.search(r"sas\s*hba", c))
        eth_part = " ".join(c for c in clauses if not re.search(fc_re, c)) if fc_part else L

        # FC HBA
        if fc_part:
            m = re.search(NUM + r"\s*g(?:b|bps|fc)?\b", fc_part)
            if m:
                found.append(_req("fc_speed_gb", ">=", float(m.group(1)), line))
            n = _ports(fc_part)
            if n:
                found.append(_req("fc_ports", ">=", n, line))

        # NIC (Ethernet)
        if (re.search(r"nic|gbe|ethernet|이더넷|네트워크|랜카드|lan\b|sfp|nw\s*포트|인터페이스", eth_part) and not re.search(r"ocp", eth_part)
                or re.search(r"\d+\s*gbe", eth_part)) and not (fc_part and not re.search(r"\d", eth_part)):
            speed_re = NUM + r"\s*(gbe|gb\s*ethernet|g\s*bps|gbps|g\b)"
            eth_clauses = [c for c in clauses if not re.search(fc_re, c)] if fc_part else clauses
            speed_clauses = [c for c in eth_clauses if re.search(speed_re, c)]
            m = re.search(speed_re, eth_part)
            n = _ports(eth_part)
            if len(speed_clauses) > 1:
                # "1GbE 4포트 / 10GbE 2포트": 속도별 포트 수를 따로 (포트 수는 그 속도 이상 포트로 검증)
                for clause in speed_clauses:
                    sp = float(re.search(speed_re, clause).group(1))
                    found.append(_req("nic_speed_gb", ">=", sp, line))
                    cn = _ports(clause)
                    if cn:
                        found.append(_req("nic_ports", ">=", cn, line, note=f"{sp:g}GbE 이상 포트", at_speed=sp))
            else:
                if m:
                    found.append(_req("nic_speed_gb", ">=", float(m.group(1)), line))
                if n:
                    found.append(_req("nic_ports", ">=", n, line))
            if not n and not speed_clauses and re.search(r"port|포트", eth_part):
                reqs.append({
                    "id": uuid.uuid4().hex[:8], "key": "manual", "label": "NIC Port 수",
                    "op": "?", "value": "", "unit": "", "source": line[:200],
                    "status": "review", "note": "문서에서 포트 수량 기준을 확인하세요",
                })

        # OCP
        if re.search(r"ocp", L):
            found.append(_req("ocp_required", "=", True, line))
            m = re.search(NUM + r"\s*gbe?", L)
            if m and not any(r["key"] == "nic_speed_gb" for r in found):
                found.append(_req("nic_speed_gb", ">=", float(m.group(1)), line, note="OCP NIC 속도"))

        # 디스크: 개수·한 개 용량·총 용량 (메모리 줄 제외)
        if re.search(r"ssd|hdd|nvme|디스크|disk|드라이브|drive|스토리지|storage|저장\s*장치|내장\s*저장", L) \
                and not re.search(r"memory|메모리|\bram\b|dimm|m\.2|boss|캐시|cache", L):
            boot = "Boot" if re.search(r"boot|부트|\bos(?![a-z])|운영\s*체제", L) else ""
            total = re.search(r"(?:총|total|합계|전체|가용|usable|실\s*용량)\s*(?:용량)?\s*[:(]?\s*" + NUM + r"\s*(tb|gb)", L)
            size = re.search(NUM + r"\s*(tb|gb)(?!\s*(?:ram|dimm|ddr))", L)
            count = re.search(r"(?<![\d.])(\d{1,3})\s*(?:개|ea|본|drives?|disks?|bays?|베이|장)\b|[x×*]\s*(\d{1,3})\b(?!\s*(?:tb|gb))", L)
            if total:
                usable = bool(re.search(r"가용|usable|실\s*용량", L))
                found.append(_req("disk_total_gb", ">=", float(total.group(1)) * (1000 if total.group(2) == "tb" else 1), line,
                                  note="가용 용량(RAID 후) — 확인 필요" if usable else boot, status="review" if usable else "auto"))
            elif size:
                found.append(_req("disk_size_gb", ">=", float(size.group(1)) * (1000 if size.group(2) == "tb" else 1), line, note=boot))
            if count:
                found.append(_req("disk_count", ">=" if re.search(GE, L) or not re.search(r"정확히|only", L) else "=", int(count.group(1) or count.group(2)), line, note=boot))

        # 폼팩터: 랙형
        if re.search(r"rack\s*(?:type|mount|형)?|랙\s*(?:형|타입|마운트)", L) and re.search(r"형태|폼\s*팩터|form|type|타입|형\b|rack\s*type|랙형", L) \
                and not re.search(r"kit|키트|rail|레일", L):
            found.append(_req("rack_mount", "=", True, line))
        # RAID 컨트롤러 (수준 표기 없이 '지원/필요')
        if re.search(r"raid\s*(?:controller|컨트롤러|카드|card)|\b(?:perc|praid)\b|하드웨어\s*raid|hw\s*raid", L):
            found.append(_req("raid_controller", "=", True, line))

        # RAID
        m = re.search(r"raid\s*-?\s*(10|1|5|6|0)\b", L)
        if m:
            note = "Boot" if re.search(r"boot|부트|os", L) else ""
            found.append(_req("raid_level", "=", f"RAID{m.group(1)}", line, note=note))

        # PSU
        if re.search(r"psu|power supply|전원|파워|power", L):
            if re.search(r"이중화|redundan|dual|이중|1\s*\+\s*1|2\s*(ea|개|unit|식)|hot.?plug|핫\s*플러그", L):
                found.append(_req("dual_psu", "=", True, line))
            w = re.search(NUM + r"\s*w\b", L)
            if w:
                found.append(_req("psu_watt", _op(L), float(w.group(1)), line))

        # Free PCIe slots
        if re.search(r"pcie|pci-e|slot|슬롯", L) and re.search(r"여유|확보|free|spare|빈|잔여|추가|향후", L):
            m = re.search(NUM + r"\s*(개|ea|slot|슬롯)", L)
            if m:
                found.append(_req("free_pcie", ">=", int(float(m.group(1))), line))

        # GPU
        if re.search(r"gpu|그래픽|가속기|accelerator", L):
            m = re.search(NUM + r"\s*(개|ea|장|way)", L)
            found.append(_req("gpu_count", _op(L) if m else ">=", int(float(m.group(1))) if m else 1, line,
                              note="" if m else "수량 미기재 → 1로 가정"))

        # 모호 표현 → 확인 필요
        if VAGUE.search(line):
            if found:
                for r in found:
                    r["status"] = "review"
                    r["note"] = (r["note"] + " / " if r["note"] else "") + "모호 표현 포함"
            elif re.search(r"서버|server|확장|성능|구성|지원|메모리|cpu|nic|디스크|disk|전원", L):
                reqs.append({"id": uuid.uuid4().hex[:8], "key": "manual", "label": "수기 검토",
                             "op": "?", "value": "", "unit": "", "source": line[:200],
                             "status": "review", "note": "정량 기준 없음 → 자동 합격 처리 금지"})
        reqs += found

    # 정량 기준이 비현실적으로 큰 경우 자동 판정 대신 사용자 확인이 필요하다.
    plausible_ranges = {
        "memory_gb": (1, 65536), "cpu_sockets": (1, 8),
        "nic_speed_gb": (1, 800), "nic_ports": (1, 128),
        "fc_speed_gb": (1, 256), "fc_ports": (1, 64),
        "psu_watt": (100, 10000), "free_pcie": (0, 64), "gpu_count": (0, 32),
        "cpu_cores": (2, 1024), "disk_count": (1, 64), "disk_size_gb": (100, 100000), "disk_total_gb": (100, 2000000),
    }
    for req in reqs:
        limits = plausible_ranges.get(req["key"])
        if limits and isinstance(req["value"], (int, float)) and not limits[0] <= req["value"] <= limits[1]:
            req["status"] = "review"
            req["note"] = (req["note"] + " / " if req["note"] else "") + "추출값이 일반적인 범위를 벗어남 — 원문 확인 필요"

    # 중복 기준은 한 건으로 합치되, 원문 근거와 확인 필요 상태는 보존한다.
    seen, out = {}, []
    for r in reqs:
        sig = (r["key"], r["op"], str(r["value"]))
        previous = seen.get(sig)
        if previous is None:
            seen[sig] = r
            out.append(r)
        else:
            if r["status"] == "review":
                previous["status"] = "review"
            sources = previous.setdefault("sources", [previous["source"]])
            if r["source"] not in sources:
                sources.append(r["source"])
            if r.get("note") and r["note"] not in previous.get("note", ""):
                previous["note"] = (previous.get("note", "") + " / " + r["note"]).strip(" /")
    return out


def _server_name(line: str) -> str | None:
    """문서의 명시적 서버/장비 제목에서 서버 그룹 이름을 추출한다."""
    value = line.strip(" \t-•*#[]()")
    value = re.sub(r"^\d+(?:[.)-]\d+)*[.)]?\s*", "", value)
    value = re.sub(r"^(?:서버|server|장비|시스템)\s*(?:명|구분|유형)?\s*[:：-]\s*", "", value, flags=re.I)
    value = re.sub(r"\s*(?:요구\s*사항|요구사양|사양서|구성\s*요구사항|requirements?)\s*$", "", value, flags=re.I)
    if not value or len(value) > 60:
        return None
    if not re.search(r"서버|server|장비|시스템", value, re.I):
        return None
    if re.search(r"cpu|memory|메모리|nic|raid|psu|disk|디스크|요약|전체", value, re.I):
        return None
    name = re.sub(r"\s*(?:서버|server|장비|시스템)\s*$", "", value, flags=re.I).strip(" -:")
    if not name:
        return None
    if re.search(r"server$", value, re.I):
        name += " 서버"
    return name[:60]


def _table_server_groups(lines: list[str]) -> list[dict] | None:
    """서버명/장비구분 열을 가진 표의 행을 서버별 요구사항으로 묶는다."""
    for header_index, line in enumerate(lines):
        headers = [cell.strip() for cell in line.split("|")]
        column = next((
            index for index, cell in enumerate(headers)
            if (re.search(r"서버|server|장비|시스템", cell, re.I)
                and re.search(r"명|구분|유형|대상|name|type|model", cell, re.I))
            or cell.strip().lower() in {"서버", "server", "장비", "시스템"}
        ), None)
        if column is None:
            continue
        grouped: dict[str, list[str]] = {}
        for row in lines[header_index + 1:]:
            cells = [cell.strip() for cell in row.split("|")]
            if len(cells) <= column:
                continue
            name = cells[column]
            if not name or name.lower() in {"-", "동일", "same"}:
                continue
            if not re.search(r"서버|server|장비|시스템", name, re.I):
                name += " 서버"
            content = " | ".join(cell for index, cell in enumerate(cells) if index != column and cell)
            if content:
                grouped.setdefault(name[:60], []).append(content)
        if len(grouped) > 1:
            return [{"name": name, "text": "\n".join(rows)} for name, rows in grouped.items()]
    return None


def extract_server_groups(text: str) -> list[dict]:
    """명시적인 서버 제목/표 열이 여러 개인 문서를 서버별로 묶는다."""
    normalized = normalize(text)
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    table_groups = _table_server_groups(lines)
    if table_groups:
        groups = table_groups
    else:
        headings: list[tuple[int, str]] = []
        for index, line in enumerate(lines):
            name = _server_name(line)
            if name and (len(line.split()) <= 8 or line.startswith(("[", "#"))):
                headings.append((index, name))
        if not headings:
            groups = [{"name": "서버 1", "text": normalized}]
        else:
            groups = []
            for heading_index, (start, name) in enumerate(headings):
                end = headings[heading_index + 1][0] if heading_index + 1 < len(headings) else len(lines)
                section = "\n".join(lines[start + 1:end]).strip()
                if section:
                    groups.append({"name": name, "text": section})
            if not groups:
                groups = [{"name": headings[0][1], "text": normalized}]

    unique: list[dict] = []
    seen_names: dict[str, int] = {}
    for index, group in enumerate(groups, 1):
        base_name = group["name"]
        seen_names[base_name] = seen_names.get(base_name, 0) + 1
        name = base_name if seen_names[base_name] == 1 else f"{base_name} {seen_names[base_name]}"
        group_text = group["text"]
        unique.append({
            "id": f"server-{index}",
            "name": name,
            "requirements": extract_requirements(group_text),
            "spec": spec_summary(group_text),
            "text": group_text,
        })
    return unique


# ------------------------------------------------------------------ 요구 사양 요약
SPEC_CATS = [
    ("CPU", r"\bcpu\b|프로세서|processor|xeon|epyc|중앙\s*처리|코어|\bcore"),
    ("Memory", r"memory|메모리|\bram\b|\bdimm|ddr[45]|주\s*기억|기억\s*장치"),
    ("Disk", r"disk|디스크|\bhdd\b|\bssd\b|nvme|저장\s*장치|스토리지|storage|하드\s*디스크|boss|m\.2"),
    ("RAID", r"\braid|레이드"),
    ("NIC", r"\bnic\b|ethernet|이더넷|\d\s*gbe|네트워크|랜\s*카드|\blan\b|sfp|10g|25g|100g|base-?t"),
    ("FC HBA", r"\bfc\b|fibre|fiber channel|\bhba\b"),
    ("OCP", r"\bocp"),
    ("PSU", r"\bpsu\b|전원|power supply|파워|\d{3,4}\s*w\b"),
    ("GPU", r"\bgpu\b|그래픽|가속기|accelerator"),
    ("PCIe/확장", r"pcie|pci-e|확장\s*슬롯|slot|슬롯|라이저|riser"),
    ("폼팩터", r"\b[12]u\b|랙\s*형|rack|폼\s*팩터|섀시|chassis"),
]


def spec_summary(text: str) -> list[dict]:
    """문서에서 카테고리별 사양 문장을 찾아 '항목: 값' 형태로 정리."""
    lines = [l.strip(" -•*·□■○●▶▷\t") for l in normalize(text).splitlines()]
    out: dict[str, list] = {}
    seen, used = set(), set()
    for i, line in enumerate(lines):
        if i in used or len(line) < 2 or len(line) > 220:
            continue
        L = line.lower()
        for cat, pat in SPEC_CATS:
            if not re.search(pat, L):
                continue
            # FC 문장이 NIC로, NIC/OCP 문장이 PSU 등으로 중복 분류되지 않도록 우선순위 처리
            if cat == "NIC" and re.search(r"\bfc\b|fibre|\bhba\b", L): continue
            if cat == "PCIe/확장" and re.search(r"\bfc\b|\bhba\b|\bnic\b|gpu", L) and not re.search(r"여유|확보|추가|free|spare", L): continue
            value = line
            # '항목명' 만 있는 짧은 줄이면 다음 줄(값)을 붙임
            if len(line) < 25 and not re.search(r"\d", line) and i + 1 < len(lines) and re.search(r"\d", lines[i + 1]):
                value = f"{line}: {lines[i + 1]}"; used.add(i + 1)
            m = re.split(r"\s*[:：]\s*", value, maxsplit=1)
            label, val = (m[0], m[1]) if len(m) == 2 and len(m[0]) < 30 else ("", value)
            key = (cat, val)
            if key in seen or not val.strip():
                continue
            seen.add(key)
            out.setdefault(cat, []).append({"label": label, "value": val.strip(), "source": line})
            break
    return [{"category": c, "items": out[c][:6]} for c, _ in SPEC_CATS if c in out]
