"""문서 구조 해석 (양식 무관).

1) 표 찾기   : 엑셀 시트 / 워드 표 / CSV / PDF·텍스트 줄을 '행×셀' 표로 통일
2) 열 역할   : 머리글 동의어 + 열 내용 형태(정수·금액·코드·문장)로 품번/품명/수량/가격/구분 판정
3) 행 분류   : 품목 / 제목 / 합계 / 메타(수신·일자) / 비고
4) 서버 그룹 : 시트 경계·표 제목·섹션 행·본체 행·목록표(서버명+수량) 신호를 점수로 조합
5) 문서 역할 : 가격 열이 있으면 '제안 구성(견적)', 아니면 '요구사항'

특정 업체 양식에 맞춘 분기는 두지 않는다. 판단 근거(evidence)와 신뢰도(confidence)를 함께 돌려주고,
애매하면 사람이 확인하도록 표시한다.
"""
from __future__ import annotations
import csv, io, re
from dataclasses import dataclass, field
from pathlib import Path

from . import parts

# ------------------------------------------------------------------ 열 머리글 동의어
ROLE_SYNONYMS = {
    "code":  r"품\s*번|부품\s*(?:번호|코드)|part\s*(?:no|number|#)|p/?n\b|sku|품목\s*코드|모델\s*코드|코드|code|order\s*no",
    "desc":  r"품\s*명|상세\s*(?:스펙|사양)|사양|규격|description|desc|spec|내역|제안\s*모델|모델명|품목\s*명",
    "desc_weak": r"품목|항목|모델|model|item",
    "qty":   r"수\s*량|qty|q'?ty|quantity|ea\b|개수|대수",
    "price": r"단\s*가|금\s*액|가\s*격|소비자|견적가?|공급가|합계|price|amount|cost|total|원\b",
    "group": r"구\s*분|용\s*도|서버\s*(?:명|구분|유형)?|장비\s*(?:명|구분)|시스템|role|server|type",
    "no":    r"^no\.?$|^번호$|^순번$|^#$|^item\s*no",
    "unit":  r"^단위$|^unit$",
}
TOTAL_ROW = re.compile(r"합\s*계|소\s*계|총\s*계|부가세|vat|sub\s*-?total|grand\s*total|^total$", re.I)
META_ROW = re.compile(r"^(수\s*신|참\s*조|발\s*신|견적\s*(?:건명|금액|일자|번호)|납기|결제|유효\s*기간|담당|연락처|이메일|tel|fax|e-?mail|date|to|from|attn)\s*[:：]?", re.I)
REMARK_ROW = re.compile(r"^(remarks?|비\s*고|참\s*고|특기\s*사항|※|■|\*)", re.I)
ROLE_WORDS = re.compile(r"\b(db|was|web|ap|app|api|backup|백업|관리|mgmt|운영|개발|테스트|test|dev|스토리지|storage|gpu|ai|로그|log|보안|파일|file|vdi|"
                        r"계산|연산|분석|배치|batch|mail|메일|dns|ldap|ad|nas|san)\b|서버|server|노드|node", re.I)


@dataclass
class Table:
    source: str                          # "시트 1", "표 2", "PDF p.3"
    rows: list[tuple[int, list[str]]]    # (원본 행번호, 셀)
    titles: list[str] = field(default_factory=list)  # 표 위쪽 제목 후보


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def load_tables(filename: str, data: bytes) -> list[Table]:
    ext = Path(filename).suffix.lower()
    out: list[Table] = []
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = []
            for n, r in enumerate(ws.iter_rows(values_only=True), 1):
                cells = [_s(c) for c in r]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append((n, cells))
            out.append(Table(f"시트 '{ws.title}'", rows))
        wb.close()
    elif ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        # 본문 문단 + 표를 문서 순서대로
        body = d.element.body
        n = 0
        para_rows: list[tuple[int, list[str]]] = []
        tcount = 0
        for el in body.iterchildren():
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "p":
                txt = "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t")).strip()
                if txt:
                    n += 1; para_rows.append((n, [txt]))
            elif tag == "tbl":
                tcount += 1
                t = next(x for x in d.tables if x._tbl is el)
                rows = []
                for r in t.rows:
                    n += 1
                    cells, prev = [], None
                    for c in r.cells:            # 병합 셀 중복 제거
                        if c._tc is prev:
                            continue
                        prev = c._tc; cells.append(_s(c.text))
                    if any(cells):
                        rows.append((n, cells))
                titles = [x[1][0] for x in para_rows[-3:]]
                if para_rows:
                    out.append(Table("본문", para_rows)); para_rows = []
                out.append(Table(f"표 {tcount}", rows, titles))
        if para_rows:
            out.append(Table("본문", para_rows))
    elif ext == ".csv":
        from .extract import _decode
        rows = [(i, [_s(c) for c in r]) for i, r in enumerate(csv.reader(io.StringIO(_decode(data))), 1) if any(x.strip() for x in r)]
        out.append(Table("CSV", rows))
    else:   # pdf / txt / json → 줄 단위. ' | ', 탭, 2칸 이상 공백을 셀 구분으로 본다
        from .extract import extract_text, normalize
        text = normalize(extract_text(filename, data))
        rows = []
        for i, line in enumerate(text.splitlines(), 1):
            if not line.strip():
                continue
            cells = [c.strip() for c in re.split(r"\s*\|\s*|\t+|\s{2,}", line) if c.strip()]
            rows.append((i, cells))
        out.append(Table("본문", rows))
    return out


# ------------------------------------------------------------------ 열 역할
def _cell_kind(v: str) -> str:
    if not v:
        return "empty"
    x = v.replace(",", "")
    if re.fullmatch(r"-?\d+", x):
        n = int(x)
        return "int_small" if 0 < n <= 999 else "money" if n >= 10000 else "int"
    if re.fullmatch(r"-?\d+\.\d+", x):
        return "float"
    if re.fullmatch(r"[A-Z0-9][A-Z0-9\-_.]{4,24}", v) and re.search(r"\d", v) and re.search(r"[A-Z]", v) and " " not in v:
        return "code"
    if len(v) >= 12 and " " in v:
        return "text"
    return "word"


def _header_roles(cells: list[str]) -> dict[int, str]:
    roles = {}
    for i, c in enumerate(cells):
        lc = c.lower()
        for role in ("no", "qty", "code", "price", "group", "desc", "desc_weak", "unit"):
            if c and len(c) <= 20 and re.search(ROLE_SYNONYMS[role], lc, re.I):
                if role not in roles.values() or role == "price":
                    roles[i] = role
                break
    # '품명/Description' 같은 강한 표현이 있으면 Item/모델 같은 약한 표현은 번호 열로 본다
    if "desc" in roles.values():
        roles = {i: ("no" if r == "desc_weak" else r) for i, r in roles.items()}
    else:
        roles = {i: ("desc" if r == "desc_weak" else r) for i, r in roles.items()}
    return roles


def _verify_roles(roles: dict[int, str], rows) -> dict[int, str]:
    """머리글로 정한 역할을 실제 내용으로 검증: 품명 열이 대부분 숫자면 번호 열로 바꾸고 문장 열을 품명으로."""
    data = [c for _, c in rows[:40]]
    def kinds(j):
        col = [c[j] for c in data if j < len(c) and c[j]]
        return [_cell_kind(v) for v in col] or ["empty"]
    for j, r in list(roles.items()):
        if r == "desc":
            k = kinds(j)
            if sum(x in ("int_small", "int") for x in k) / len(k) > 0.6:
                roles[j] = "no"
                cand = [(sum(x == "text" for x in kinds(i)), i) for i in range(max(map(len, data))) if i not in roles]
                if cand and max(cand)[0]:
                    roles[max(cand)[1]] = "desc"
    return roles


def find_header(rows) -> tuple[int, dict[int, str]] | None:
    """표 머리글 행(인덱스)과 열 역할. 머리글이 없으면 열 내용 형태로 추정."""
    best = None
    for idx, (_, cells) in enumerate(rows):
        roles = _header_roles(cells)
        kinds = set(roles.values())
        score = len(kinds & {"code", "desc", "qty"}) * 2 + ("price" in kinds) + ("group" in kinds) + ("no" in kinds) * 0.5
        # 머리글 바로 아래 행이 실제 데이터 같아야 함
        nxt = rows[idx + 1: idx + 4]
        if score >= 3 and len(cells) >= 2 and any(any(_cell_kind(c) in ("int_small", "code", "text") for c in r[1]) for r in nxt):
            if not best or score > best[2]:
                best = (idx, roles, score)
    if best:
        return best[0], _verify_roles(best[1], rows[best[0] + 1:])
    return _infer_roles(rows)


def _infer_roles(rows):
    """머리글이 없을 때: 행 다수가 같은 열 패턴(코드/문장/작은정수)을 가지면 그 패턴으로 역할을 정한다."""
    data = [c for _, c in rows if len(c) >= 2]
    if len(data) < 3:
        return None
    width = max(len(c) for c in data)
    stats = []
    for j in range(width):
        col = [c[j] for c in data if j < len(c) and c[j]]
        kinds = [_cell_kind(v) for v in col]
        stats.append({k: kinds.count(k) / max(1, len(data)) for k in set(kinds)})
    roles = {}
    def pick(kind, role, thr=0.5):
        cand = [(s.get(kind, 0), j) for j, s in enumerate(stats) if j not in roles]
        if cand:
            v, j = max(cand)
            if v >= thr:
                roles[j] = role
    pick("text", "desc"); pick("code", "code", 0.4); pick("int_small", "qty", 0.5); pick("money", "price", 0.4)
    if "desc" in roles.values() and "qty" in roles.values():
        return -1, roles
    return None


# ------------------------------------------------------------------ 행 분류 · 블록
@dataclass
class Item:
    code: str
    desc: str
    qty: float | None
    group: str
    where: str
    interp: dict


@dataclass
class Block:
    name: str | None
    name_how: str
    items: list[Item]
    where: str
    has_price: bool
    is_summary: bool = False


def _row_text(cells):
    return " ".join(c for c in cells if c)


def _is_title(cells) -> str | None:
    vals = [c for c in cells if c]
    if not vals or len(vals) > 2:
        return None
    t = vals[0]
    if TOTAL_ROW.search(t) or META_ROW.search(t) or REMARK_ROW.search(t) or len(t) > 40:
        return None
    if len(vals) == 2 and not re.search(r"단위|unit|원\)", vals[1], re.I):
        return None
    t = re.sub(r"^[\s\-•□■○●▶▷\[\(]*(?:\d+[.)]|[가-하][.)]|[IVX]+\.)?\s*", "", t).strip(" ])")
    return t or None


def table_blocks(t: Table) -> list[Block]:
    rows = t.rows
    hdr = find_header(rows)
    if not hdr:
        return []
    hi, roles = hdr
    col = {r: j for j, r in roles.items() if r != "price"}
    has_price = "price" in roles.values()
    # 표 위쪽 제목(머리글 위 3행 안, 메타 행 제외)
    title = None
    for _, cells in reversed(rows[max(0, hi - 3):max(0, hi)]):
        tt = _is_title(cells)
        if tt and not re.search(r"견\s*적\s*서|quotation|사양서|요구사항", tt, re.I):
            title = tt; break
    title = title or next((x for x in reversed(t.titles) if ROLE_WORDS.search(x)), None)
    blocks: list[Block] = []
    cur = Block(title, "표 제목" if title else "", [], t.source, has_price)
    for n, cells in rows[hi + 1:]:
        txt = _row_text(cells)
        if TOTAL_ROW.search(cells[0] if cells else "") or TOTAL_ROW.fullmatch(txt.split(" ")[0] or ""):
            continue
        if REMARK_ROW.search(txt) or META_ROW.search(txt):
            continue
        get = lambda r: cells[col[r]] if r in col and col[r] < len(cells) else ""
        desc, code, qtys, grp = get("desc"), get("code"), get("qty"), get("group")
        qty = parts._num(qtys.replace(",", "")) if qtys else None
        sec = _is_title(cells)
        if sec and not qty and not code:
            # 섹션 행: 새 블록 시작
            if cur.items:
                blocks.append(cur)
            cur = Block(sec, "섹션 행", [], t.source, has_price)
            continue
        if not desc and not code:
            continue
        if not desc and code and _cell_kind(code) != "code":
            desc, code = code, ""
        it = Item(code, desc or code, qty, grp, f"{t.source} {n}행", parts.interpret(code, desc or code))
        # 본체 행이 이미 있는 블록에 또 본체가 나오면 새 서버 구성 시작
        if it.interp["category"] == "base" and any(i.interp["category"] == "base" for i in cur.items) and not grp:
            blocks.append(cur)
            cur = Block(None, "", [], t.source, has_price)
        cur.items.append(it)
    if cur.items:
        blocks.append(cur)
    # 목록표 판정: '구분' 열이 있고 품목 대부분이 본체(서버 모델)이면 서버 목록표
    for b in blocks:
        if "group" in col and b.items:
            base_like = sum(1 for i in b.items if i.interp["category"] == "base")
            if base_like >= max(1, len(b.items) * 0.5):
                b.is_summary = True
    return blocks


# ------------------------------------------------------------------ 그룹 조립
def _key(name: str) -> str:
    return re.sub(r"[\s_\-·]|서버|server|용", "", (name or "").lower())


def analyze(filename: str, data: bytes) -> dict:
    tables = load_tables(filename, data)
    blocks = [b for t in tables for b in table_blocks(t)]
    has_price = any(b.has_price for b in blocks)
    summary = [i for b in blocks if b.is_summary for i in b.items]
    details = [b for b in blocks if not b.is_summary and b.items]

    groups = []
    used_summary = set()
    for b in details:
        base = next((i for i in b.items if i.interp["category"] == "base"), None)
        name, how, conf = b.name, b.name_how, 0.6
        evidence = [f"{b.where}: 품목 {len(b.items)}개"]
        # 목록표의 서버명과 연결
        link = None
        for k, s in enumerate(summary):
            if k in used_summary:
                continue
            if name and (_key(s.group) == _key(name) or _key(s.group) in _key(name) or _key(name) in _key(s.group)):
                link = (k, s); break
        if link:
            used_summary.add(link[0]); s = link[1]
            qty = s.qty; conf = 0.9
            evidence.append(f"목록표 '{s.group}' ({s.where}) 수량 {s.qty:g}대와 연결")
        else:
            qty = None
        if base:
            evidence.append(f"본체 행: {base.desc} ({base.where})")
            conf = max(conf, 0.75)
        if not name:
            name = base.desc if base else None
            how = "본체 품명" if base else ""
        # 수량이 총량인지 1대 기준인지
        unit_note = None
        base_qty = base.qty if base and base.qty else None
        if qty is None and base_qty:
            qty = base_qty
            evidence.append(f"본체 수량 {base_qty:g}을 서버 대수로 사용")
        per_unit = 1.0
        if base_qty and base_qty > 1:
            others = [i.qty for i in b.items if i is not base and i.qty]
            if others and all(abs((q / base_qty) - round(q / base_qty)) < 1e-6 for q in others):
                per_unit = base_qty
                unit_note = f"상세 수량이 본체 {base_qty:g}대분 합계로 보여 대당 수량으로 환산"
            else:
                unit_note = "상세 수량이 대당인지 합계인지 판단 불가 — 확인 필요"
                conf = min(conf, 0.6)
        groups.append(_make_group(len(groups) + 1, name, how, qty, b, per_unit, evidence, conf, unit_note))

    # 목록표에만 있고 상세가 없는 서버
    for k, s in enumerate(summary):
        if k in used_summary:
            continue
        b = Block(s.group, "목록표", [s], s.where, has_price)
        groups.append(_make_group(len(groups) + 1, s.group or s.desc, "목록표", s.qty, b, 1.0,
                                  [f"목록표 행 {s.where} (상세 구성 없음)"], 0.5, "상세 구성이 문서에 없음"))

    # 서버(본체)가 없는 블록 = 공통/기타 품목
    common_items = []
    server_groups = []
    for g in groups:
        if not g["_has_base"] and len(groups) > 1 and not any(c in g["_cats"] for c in ("cpu", "memory")):
            common_items += g["items"]
        else:
            server_groups.append(g)
    for g in server_groups:
        g.pop("_has_base"); g.pop("_cats")
    return {
        "doc_role": "quote" if has_price else ("requirement" if not server_groups else "spec_table"),
        "groups": server_groups,
        "common_items": common_items,
        "tables": [{"source": t.source, "rows": len(t.rows)} for t in tables],
    }


def _make_group(idx, name, how, qty, b: Block, per_unit, evidence, conf, unit_note):
    items, cats = [], set()
    for i in b.items:
        q = (i.qty / per_unit) if (i.qty and per_unit > 1) else i.qty
        cat = i.interp["category"]
        cats.add(cat)
        items.append({"code": i.code, "desc": i.desc, "qty": q, "category": cat,
                      "category_ko": parts.CATEGORY_KO.get(cat, cat), "attrs": i.interp["attrs"],
                      "confidence": i.interp["confidence"], "how": i.interp.get("how"), "where": i.where})
    base = next((x for x in items if x["category"] == "base"), None)
    return {
        "id": f"server-{idx}",
        "name": name or f"서버 {idx}",
        "name_how": how,
        "quantity": int(qty) if qty else None,
        "model_hint": (base or {}).get("attrs", {}).get("model"),
        "base_desc": (base or {}).get("desc"),
        "items": items,
        "proposed": proposal(items),
        "evidence": evidence,
        "confidence": round(conf, 2),
        "notes": [n for n in [unit_note] if n] + [f"해석 안 된 품목 {sum(1 for x in items if x['category']=='unknown')}개"
                                                  for _ in [0] if any(x["category"] == "unknown" for x in items)],
        "_has_base": base is not None,
        "_cats": cats,
    }


# ------------------------------------------------------------------ 제안 구성 요약
def proposal(items: list[dict]) -> dict:
    """품목 → 서버 구성 요약(대당). 가격은 쓰지 않는다."""
    by = lambda c: [x for x in items if x["category"] == c]
    q = lambda x: x["qty"] or 1
    cpu = by("cpu"); mem = by("memory"); drv = by("drive"); psu = by("psu")
    mem_total = sum((x["attrs"].get("size_gb") or 0) * q(x) for x in mem)
    dimms = {}
    for x in mem:
        s = x["attrs"].get("size_gb")
        if s:
            dimms[s] = dimms.get(s, 0) + q(x)
    return {
        "cpu": {"model": cpu[0]["attrs"].get("model") if cpu else None,
                "count": int(sum(q(x) for x in cpu)) if cpu else 0,
                "cores": cpu[0]["attrs"].get("cores") if cpu else None},
        "memory": {"total_gb": mem_total, "dimms": [{"size_gb": k, "qty": int(v)} for k, v in sorted(dimms.items())]},
        "drives": [{"desc": x["desc"], "qty": int(q(x)), **x["attrs"]} for x in drv],
        "raid": [x["desc"] for x in by("raid")],
        "nic": [{"desc": x["desc"], "qty": int(q(x)), **x["attrs"]} for x in by("nic")],
        "ocp": [{"desc": x["desc"], "qty": int(q(x)), **x["attrs"]} for x in by("ocp")],
        "fc": [{"desc": x["desc"], "qty": int(q(x)), **x["attrs"]} for x in by("fc")],
        "gpu": [{"desc": x["desc"], "qty": int(q(x))} for x in by("gpu")],
        "riser": [{"desc": x["desc"], "qty": int(q(x)), **x["attrs"]} for x in by("riser")],
        "psu": {"watt": psu[0]["attrs"].get("watt") if psu else None, "count": int(sum(q(x) for x in psu)) if psu else 0},
    }


# ------------------------------------------------------------------ 열 단위 서버 표 (사양 비교표)
def column_groups(t: Table) -> list[dict] | None:
    """'구분 | DB서버 | WAS서버' 처럼 서버가 열로 나열된 표 → 서버별 '항목: 값' 문장."""
    for idx, (_, cells) in enumerate(t.rows[:40]):
        heads = [(j, c) for j, c in enumerate(cells) if j > 0 and c and len(c) <= 30 and ROLE_WORDS.search(c)
                 and not re.search(r"cpu|memory|메모리|disk|nic|수량|qty|단가|금액", c, re.I)]
        if len(heads) < 2:
            continue
        groups = {j: {"name": c, "lines": [], "qty": None, "where": f"{t.source} {t.rows[idx][0]}행 머리글"} for j, c in heads}
        for n, row in t.rows[idx + 1:]:
            label = row[0] if row else ""
            if not label or TOTAL_ROW.search(label):
                continue
            for j, g in groups.items():
                v = row[j] if j < len(row) else ""
                if not v or v in ("-", "동일", "상동"):
                    continue
                if re.search(r"수\s*량|대\s*수|qty|quantity", label, re.I):
                    m = re.search(r"(\d+)", v); g["qty"] = int(m.group(1)) if m else None
                g["lines"].append((f"{label}: {v}", f"{t.source} {n}행"))
        out = [g for g in groups.values() if g["lines"]]
        return out if len(out) >= 2 else None
    return None


HEAD_QTY = re.compile(r"(\d+)\s*(?:대|식|set|ea|units?|nodes?)\b", re.I)


def clean_name(name: str) -> str:
    n = re.sub(r"^[\s\-•□■○●▶▷\[\(]*(?:\d+(?:\.\d+)*[.)]|[가-하][.)]|[IVX]+\.)?\s*", "", name or "")
    n = re.sub(r"\s*[\(\[]?\s*\d+\s*(?:대|식|set|ea|units?|nodes?)\s*[\)\]]?\s*$", "", n, flags=re.I)
    return n.strip(" -:()[]") or name


def analyze_document(filename: str, data: bytes, text: str) -> dict:
    """업로드 문서 → 통합 그룹 구조. 견적서는 제안 구성, 그 외는 요구사항."""
    from . import extract
    res = analyze(filename, data)
    groups = []
    if res["doc_role"] == "quote" and res["groups"]:
        for g in res["groups"]:
            spec = {}
            for it in g["items"]:
                if it["category"] in ("accessory",):
                    continue
                q = f" × {it['qty']:g}" if it["qty"] else ""
                spec.setdefault(it["category_ko"], []).append({
                    "label": it["code"], "value": f"{it['desc']}{q}", "source": it["where"],
                    "kind": "quote", "confidence": it["confidence"]})
            groups.append({**g, "doc_role": "quote", "requirements": [],
                           "spec": [{"category": k, "items": v} for k, v in spec.items()]})
        return {"doc_role": "quote", "groups": groups, "common_items": res["common_items"], "tables": res["tables"]}

    # 요구사항 문서: 열 단위 서버 표 → 그 외 제목 기반 분리(기존 규칙)
    for t in load_tables(filename, data):
        cg = column_groups(t)
        if cg:
            for i, g in enumerate(cg, 1):
                body = "\n".join(l for l, _ in g["lines"])
                groups.append({"id": f"server-{i}", "name": g["name"], "quantity": g["qty"], "doc_role": "requirement",
                               "requirements": extract.extract_requirements(body), "spec": extract.spec_summary(body),
                               "evidence": [f"서버가 열로 나열된 표 ({g['where']})"], "confidence": 0.85, "notes": []})
            return {"doc_role": "requirement", "groups": groups, "common_items": [], "tables": res["tables"]}
    rule = extract.extract_server_groups(text)
    for g in rule:
        m = HEAD_QTY.search(g["name"]) or None
        groups.append({**g, "name": clean_name(g["name"]), "quantity": g.get("quantity") or (int(m.group(1)) if m else None), "doc_role": "requirement",
                       "evidence": ["문서 제목/구간 기준 분리" if len(rule) > 1 else "서버 구분 없음 — 문서 전체를 1개 서버로 처리"],
                       "confidence": 0.7 if len(rule) > 1 else 0.5,
                       "notes": [] if len(rule) > 1 else ["서버가 여러 대라면 '서버 나누기'로 구간을 지정하세요"]})
    return {"doc_role": "requirement", "groups": groups, "common_items": [], "tables": res["tables"]}
